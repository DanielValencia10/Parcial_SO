"""
reporte.py — arma tablas y graficas comparativas a partir de resultados/mediciones.csv

Genera resultados/comparacion.xlsx con:
  - Comparacion: por cada tamano de matriz, tabla con tiempo promedio de V1..V4, speedup contra V1,
                 eficiencia (speedup / hilos) y la version mas rapida, + grafica de tiempos y de speedup
  - Tamano:      tiempo vs tamano de la matriz (una tabla y grafica por cantidad de hilos/procesos)
  - Equipo:      datos de resultados/equipo.json

Uso:  python reporte.py
      python reporte.py --entrada resultados/mediciones.csv --iteraciones 200
Requiere:  pip install openpyxl
"""

import argparse
import csv
import json
import os
import statistics
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

CARPETA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resultados")
VERS = [1, 2, 3, 4]
NOMBRES = {1: "V1 Secuencial", 2: "V2 Hilos", 3: "V3 Procesos", 4: "V4 NumPy"}
COLORES = {1: "96AABE", 2: "5694FF", 3: "FF9942", 4: "5CCE78"}   # mismos colores de la GUI
GRIS = "7F7F7F"

_lado = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def cabecera(celda):
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor="1F3A5F")
    celda.alignment, celda.border = CENTRO, BORDE


def celda_dato(ws, fila, col, valor, formato=None):
    c = ws.cell(row=fila, column=col, value=valor)
    c.alignment, c.border = CENTRO, BORDE
    if formato:
        c.number_format = formato
    return c


# --- lectura
def leer(ruta, iteraciones=None):
    if not os.path.exists(ruta):
        raise SystemExit(f"No existe {ruta}: primero hay que medir (python medicion.py).")
    with open(ruta, newline="", encoding="utf-8") as f:
        filas = list(csv.DictReader(f))
    if not filas:
        raise SystemExit(f"{ruta} esta vacio.")
    its = iteraciones or max(int(f["iteraciones"]) for f in filas)
    datos = defaultdict(list)          # (version, tamano, trabajadores) -> [segundos, ...]
    for f in filas:
        if int(f["iteraciones"]) == its:
            datos[(int(f["version"]), int(f["tamano"]), int(f["trabajadores"]))].append(float(f["segundos"]))
    if not datos:
        raise SystemExit(f"No hay mediciones con {its} iteraciones en {ruta}.")
    return datos, its


# --- estilo de graficas
def estilo_grafica(g, titulo, x, y, ancho=15, alto=7.5, log=True, formato_y=None):
    g.title, g.x_axis.title, g.y_axis.title = titulo, x, y
    g.width, g.height = ancho, alto
    g.x_axis.delete = False            # sin esto algunas versiones de Excel ocultan los ejes
    g.y_axis.delete = False
    g.legend.position = "b"
    if log:
        g.y_axis.scaling.logBase = 10  # V4 es mucho mas rapida que V1: en escala lineal no se veria
    if formato_y:
        g.y_axis.number_format = formato_y


def pintar_barras(g, colores):
    for serie, color in zip(g.series, colores):
        serie.graphicalProperties = GraphicalProperties(solidFill=color)
        serie.graphicalProperties.line.solidFill = color


def pintar_lineas(g, colores, ultima_punteada=False):
    for i, (serie, color) in enumerate(zip(g.series, colores)):
        base = ultima_punteada and i == len(colores) - 1
        serie.graphicalProperties = GraphicalProperties()
        serie.graphicalProperties.line.solidFill = color
        serie.graphicalProperties.line.width = 22000 if base else 28000
        serie.smooth = False
        if base:
            serie.graphicalProperties.line.dashStyle = "dash"
            serie.marker = Marker(symbol="none")
        else:
            serie.marker = Marker(symbol="circle", size=7)
            serie.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            serie.marker.graphicalProperties.line.solidFill = color


# --- hojas
def hoja_comparacion(wb, datos, its, tamanos, trabajadores, t):
    ws = wb.active
    ws.title = "Comparacion"
    ws["A1"] = f"Comparacion de versiones — promedio de las ejecuciones, {its} iteraciones"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Speedup = tiempo V1 / tiempo de la version (mayor que 1x = mas rapida que V1). "
                "Eficiencia = speedup / hilos (solo V2 y V3). V1 y V4 usan 1 solo hilo.")
    titulos = ["Hilos/proc", *[NOMBRES[v] for v in VERS], "Speedup V2", "Speedup V3", "Speedup V4",
               "Eficiencia V2", "Eficiencia V3", "Mas rapido", "Ref. V1 (1x)"]
    formatos = {2: '0.000" s"', 3: '0.000" s"', 4: '0.000" s"', 5: '0.000" s"',
                6: '0.00"x"', 7: '0.00"x"', 8: '0.00"x"', 9: "0.0%", 10: "0.0%", 12: '0.00"x"'}

    fila = 4
    for n in tamanos:
        ws.cell(row=fila, column=1, value=f"Matriz {n} x {n}").font = Font(bold=True, size=12)
        cab = fila + 1
        for i, texto in enumerate(titulos, 1):
            cabecera(ws.cell(row=cab, column=i, value=texto))
        ws.row_dimensions[cab].height = 32

        for i, p in enumerate(trabajadores):
            r = cab + 1 + i
            ts = {v: t(v, n, p) for v in VERS}
            sp = {v: (ts[1] / ts[v] if ts[1] and ts[v] else None) for v in (2, 3, 4)}
            disp = {v: x for v, x in ts.items() if x}
            valores = [p, ts[1], ts[2], ts[3], ts[4], sp[2], sp[3], sp[4],
                       sp[2] / p if sp[2] else None, sp[3] / p if sp[3] else None,
                       NOMBRES[min(disp, key=disp.get)] if disp else None, 1]
            for c, val in enumerate(valores, 1):
                celda_dato(ws, r, c, val, formatos.get(c))

        fin = cab + len(trabajadores)
        cats = Reference(ws, min_col=1, min_row=cab + 1, max_row=fin)

        g = BarChart()
        g.type, g.grouping = "col", "clustered"
        g.add_data(Reference(ws, min_col=2, max_col=5, min_row=cab, max_row=fin), titles_from_data=True)
        g.set_categories(cats)
        estilo_grafica(g, f"Tiempo por version — matriz {n}x{n} (escala log)",
                       "Hilos / procesos", "Tiempo (s, escala log)")
        pintar_barras(g, [COLORES[v] for v in VERS])
        ws.add_chart(g, f"N{fila}")

        s = LineChart()
        s.add_data(Reference(ws, min_col=6, max_col=8, min_row=cab, max_row=fin), titles_from_data=True)
        s.add_data(Reference(ws, min_col=12, max_col=12, min_row=cab, max_row=fin), titles_from_data=True)
        s.set_categories(cats)
        estilo_grafica(s, f"Speedup contra V1 — matriz {n}x{n} (escala log)",
                       "Hilos / procesos", "Speedup (x, escala log)", formato_y='0.0"x"')
        pintar_lineas(s, [COLORES[2], COLORES[3], COLORES[4], GRIS], ultima_punteada=True)
        ws.add_chart(s, f"W{fila}")

        fila += max(len(trabajadores) + 4, 17)

    for i, ancho in enumerate([12, 15, 13, 13, 13, 12, 12, 12, 13, 13, 16, 13], 1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    ws.freeze_panes = "A4"


def hoja_tamano(wb, tamanos, trabajadores, t):
    ws = wb.create_sheet("Tamano")
    ws["A1"] = "Tiempo vs tamano de la matriz"
    ws["A1"].font = Font(bold=True, size=13)
    fila = 3
    for p in trabajadores:
        ws.cell(row=fila, column=1, value=f"V2 y V3 con {p} hilos/procesos").font = Font(bold=True, size=12)
        cab = fila + 1
        for i, texto in enumerate(["Matriz", *[NOMBRES[v] for v in VERS]], 1):
            cabecera(ws.cell(row=cab, column=i, value=texto))
        ws.row_dimensions[cab].height = 32
        for i, n in enumerate(tamanos):
            r = cab + 1 + i
            celda_dato(ws, r, 1, f"{n} x {n}")
            for c, v in enumerate(VERS, 2):
                celda_dato(ws, r, c, t(v, n, p), '0.000" s"')
        fin = cab + len(tamanos)

        g = LineChart()
        g.add_data(Reference(ws, min_col=2, max_col=5, min_row=cab, max_row=fin), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=cab + 1, max_row=fin))
        estilo_grafica(g, f"Tiempo vs tamano de la matriz — {p} hilos/procesos (escala log)",
                       "Tamano de la matriz", "Tiempo (s, escala log)")
        pintar_lineas(g, [COLORES[v] for v in VERS])
        ws.add_chart(g, f"G{fila}")
        fila += max(len(tamanos) + 4, 17)

    for i, ancho in enumerate([14, 15, 13, 13, 13], 1):
        ws.column_dimensions[chr(64 + i)].width = ancho


def hoja_equipo(wb, carpeta):
    ruta = os.path.join(carpeta, "equipo.json")
    if not os.path.exists(ruta):
        return
    with open(ruta, encoding="utf-8") as f:
        equipo = json.load(f)
    ws = wb.create_sheet("Equipo")
    for i, (k, v) in enumerate(equipo.items(), 1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=str(v))
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 70


# --- principal
def generar(ruta_csv, salida, iteraciones=None):
    datos, its = leer(ruta_csv, iteraciones)
    tamanos = sorted({n for (_, n, _) in datos})
    trabajadores = sorted({p for (v, _, p) in datos if v in (2, 3)}) or [1]

    def t(v, n, p):
        """Promedio de las ejecuciones. V1 y V4 se midieron con 1 solo hilo."""
        k = (v, n, 1 if v in (1, 4) else p)
        return statistics.mean(datos[k]) if k in datos else None

    if not any(t(1, n, 1) for n in tamanos):
        raise SystemExit("Faltan mediciones de V1: el speedup se calcula contra V1. Corre con --versiones 1,2,3,4.")

    wb = Workbook()
    hoja_comparacion(wb, datos, its, tamanos, trabajadores, t)
    hoja_tamano(wb, tamanos, trabajadores, t)
    hoja_equipo(wb, salida)

    ruta_xlsx = os.path.join(salida, "comparacion.xlsx")
    try:
        wb.save(ruta_xlsx)
        print(f"Tablas y graficas guardadas en: {ruta_xlsx}")
    except PermissionError:
        print(f"No se pudo guardar {ruta_xlsx}: esta abierto en Excel. Cierralo y corre:  python reporte.py")


def main():
    ap = argparse.ArgumentParser(description="Tablas y graficas comparativas desde mediciones.csv")
    ap.add_argument("--entrada", default=os.path.join(CARPETA, "mediciones.csv"))
    ap.add_argument("--salida", default=CARPETA)
    ap.add_argument("--iteraciones", type=int, default=None, help="por defecto, el mayor valor del CSV")
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    generar(a.entrada, a.salida, a.iteraciones)


if __name__ == "__main__":
    main()